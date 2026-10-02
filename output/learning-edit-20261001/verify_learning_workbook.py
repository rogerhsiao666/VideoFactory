"""Verify every exported cell and render legible column sections for visual QA."""

import argparse
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from PIL import Image, ImageDraw, ImageFont

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
import cards
import learning_editor as editor


def wrap(text, width, font):
    result = []
    for paragraph in text.split("\n"):
        remaining = paragraph
        while remaining:
            index = 1
            while index <= len(remaining) and font.getlength(remaining[:index]) <= width:
                index += 1
            index = max(1, index - 1)
            if index < len(remaining):
                space = remaining.rfind(" ", 0, index + 1)
                if space > index // 3:
                    index = space
            result.append(remaining[:index].strip())
            remaining = remaining[index:].lstrip()
    return result or [""]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    parser.add_argument("--source", required=True)
    parser.add_argument("--report-dir", type=Path,
                        default=Path(__file__).resolve().parents[2] / "temp" / "learning-preview",
                        help="Keep disposable previews and reports out of the deliverables directory")
    parser.add_argument("--editorial", action="store_true", help="核對本機編修版，不宣稱外部API審稿通過")
    args = parser.parse_args()
    target = Path(args.workbook)
    source_path = Path(args.source)
    source_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
    assert source_hash == "90e6876a0188743498304fe3a8b400f1827b2e70f8eea00f8a38f4281e78c50a"
    source = cards.load_xlsx_items(str(source_path))
    suffix = ".editor.json" if args.editorial else ".edit.json"
    state = json.loads(Path(str(target) + suffix).read_text(encoding="utf-8"))
    assert state["source"] == source and len(source) == 50
    if args.editorial:
        assert state["editorial_method"] == "Codex ESL editorial revision"
        assert state["external_api_review"] == "not_completed_permission_timeout"
        assert all(group["boundary"].strip() for group in state["groups"])
    else:
        assert state["review_passed"] is True
    editor.validate_deck(state["items"], state["groups"], source)
    semantic = defaultdict(list)
    for item in state["items"]:
        semantic[item["core"] if args.editorial else item["_semantic_group"]].append(item)
    level_key = "tier" if args.editorial else "_semantic_level"
    assert all(len(pair) == 2 and {item[level_key] for item in pair} == {"basic", "advanced"}
               for pair in semantic.values())
    for group in state["groups"]:
        pair = [item for item in state["items"] if item["core"] == group["core"]]
        if not args.editorial:
            assert len({item["_semantic_group"] for item in pair}) == 1, group["core"]
    quote_pair = [item for item in state["items"] if item["core"] == "取得書面報價與明細"]
    assert len(quote_pair) == 2
    assert all(re.search(r"written (?:quote|estimate)|(?:quote|costs) in writing", item["word_en"].lower())
               for item in quote_pair)
    workbook = openpyxl.load_workbook(target)
    assert workbook.sheetnames == ["Cards"]
    sheet = workbook.active
    rows = list(sheet.iter_rows())
    assert sheet.max_column == 12 and len(rows) == len(state["items"]) + 1
    assert [cell.value for cell in rows[0]] == editor.LEARNING_HEADERS
    for row, item in zip(rows[1:], state["items"]):
        assert [cell.value for cell in row] == [item[key] for key in editor.LEARNING_HEADERS]
    assert all(isinstance(cell.value, str) and cell.value.strip() for row in rows for cell in row)
    assert not any(cell.data_type in ("e", "f") for row in rows for cell in row)
    assert sheet.freeze_panes == "E2"
    assert sheet.auto_filter.ref == f"A1:L{len(rows)}"
    assert f"$A$1:$L${len(rows)}" in str(sheet.print_area)
    assert all(cell.fill.fgColor.rgb == "004472C4" for cell in rows[0])
    assert all(cell.alignment.wrap_text for row in rows[1:] for cell in row)
    report_directory = args.report_dir
    report_directory.mkdir(parents=True, exist_ok=True)
    widths = [round(sheet.column_dimensions[cell.column_letter].width * 7 + 5) for cell in rows[0]]
    font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Unicode.ttf", 15)
    # PIL does not apply the system font fallback that Excel uses for emoji.
    emoji_font = ImageFont.truetype("/System/Library/Fonts/Apple Color Emoji.ttc", 160)
    star = Image.new("RGBA", (160, 160))
    ImageDraw.Draw(star).text((0, 0), "⭐", font=emoji_font, embedded_color=True, anchor="lt")
    star = star.resize((16, 16), Image.Resampling.LANCZOS)
    sections = [("phrases", list(range(7))), ("actions", [0, 7, 8]), ("examples", [0, 9, 10, 11])]
    previews, clipped = [], []
    for start in range(1, len(rows), 10):
        selected = [rows[0]] + rows[start:start + 10]
        heights = [round(sheet.row_dimensions[row[0].row].height * 96 / 72) for row in selected]
        for name, indices in sections:
            image = Image.new("RGB", (sum(widths[i] for i in indices), sum(heights)), "white")
            draw = ImageDraw.Draw(image)
            y = 0
            for row, height in zip(selected, heights):
                x = 0
                for i in indices:
                    cell, width = row[i], widths[i]
                    fill, color = ("#4472C4", "white") if cell.row == 1 else ("white", "#202020")
                    draw.rectangle((x, y, x + width, y + height), fill=fill, outline="#D9D9D9")
                    if cell.column == 3 and cell.row != 1:
                        for offset in range(cell.value.count("⭐")):
                            image.paste(star, (x + 5 + offset * 18, y + 2), star)
                        x += width
                        continue
                    lines = wrap(cell.value, width - 10, font)
                    if len(lines) * 18 + 4 > height:
                        clipped.append(cell.coordinate)
                    for line_number, line in enumerate(lines):
                        draw.text((x + 5, y + 2 + line_number * 18), line, font=font, fill=color)
                    x += width
                y += height
            preview = report_directory / f"preview-{start:02d}-{name}.png"
            image.save(preview)
            previews.append(preview.name)
    workbook.close()
    report = {"source_sha256": source_hash, "source_rows": len(source), "concepts": len(state["groups"]),
              "output_rows": len(state["items"]), "columns": 12, "all_cells_match_approved_content": True,
              "review_method": state.get("editorial_method", "external_api"),
              "external_api_review": state.get("external_api_review", "passed"),
              "basic_advanced_pairs": len(semantic), "scenarios": dict(Counter(x["Scenario"] for x in state["items"])),
              "previews": previews, "clipped_cells": clipped}
    (report_directory / "verification.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))
    assert not clipped, clipped


if __name__ == "__main__":
    main()
